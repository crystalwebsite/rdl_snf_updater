from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import tempfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# -----------------------------------------------------------------------------
# Patterns
# -----------------------------------------------------------------------------
DATASET_BLOCK_RE = re.compile(
    r"<(?:(?:[A-Za-z_][\w.\-]*):)?DataSet\b(?P<attrs>[^>]*)>"
    r"(?P<body>.*?)"
    r"</(?:(?:[A-Za-z_][\w.\-]*):)?DataSet\s*>",
    re.IGNORECASE | re.DOTALL,
)

DATASET_NAME_RE = re.compile(
    r"\bName\s*=\s*(?P<quote>['\"])(?P<name>.*?)(?P=quote)",
    re.IGNORECASE | re.DOTALL,
)

COMMAND_TEXT_RE = re.compile(
    r"(?P<open><(?:(?:[A-Za-z_][\w.\-]*):)?CommandText\b[^>]*>)"
    r"(?P<cmd>.*?)"
    r"(?P<close></(?:(?:[A-Za-z_][\w.\-]*):)?CommandText\s*>)",
    re.IGNORECASE | re.DOTALL,
)

# Matches the first two positional string parameters of:
# Source = Snowflake.Databases("server", "warehouse", ...)
# It also supports XML-escaped quotes (&quot;).
SNOWFLAKE_SOURCE_RE = re.compile(
    r"(?P<head>\bSource\s*=\s*Snowflake\.Databases\s*\(\s*)"
    r"(?P<q1>\"|&quot;)(?P<server>.*?)(?P=q1)"
    r"(?P<separator>\s*,\s*)"
    r"(?P<q2>\"|&quot;)(?P<warehouse>.*?)(?P=q2)",
    re.IGNORECASE | re.DOTALL,
)


AUDIT_HEADERS = [
    "Timestamp",
    "File Name",
    "Dataset Name",
    "Status",
    "Before Server",
    "After Server",
    "Before Warehouse",
    "After Warehouse",
    "Before Source Code",
    "After Source Code",
    "Input Path",
    "Output Path",
    "Original SHA256",
    "Modified SHA256",
]


@dataclass
class AuditRow:
    timestamp: str
    file_name: str
    dataset_name: str
    status: str
    before_server: str = ""
    after_server: str = ""
    before_warehouse: str = ""
    after_warehouse: str = ""
    before_source_code: str = ""
    after_source_code: str = ""
    input_path: str = ""
    output_path: str = ""
    original_sha256: str = ""
    modified_sha256: str = ""

    def as_list(self) -> list[str]:
        return [
            self.timestamp,
            self.file_name,
            self.dataset_name,
            self.status,
            self.before_server,
            self.after_server,
            self.before_warehouse,
            self.after_warehouse,
            self.before_source_code,
            self.after_source_code,
            self.input_path,
            self.output_path,
            self.original_sha256,
            self.modified_sha256,
        ]


@dataclass
class FileProcessResult:
    input_file: Path
    output_file: Path | None
    dataset_names: list[str]
    changed_count: int
    deleted_original: bool


# -----------------------------------------------------------------------------
# Encoding / XML helpers
# -----------------------------------------------------------------------------
def decode_xml_bytes(data: bytes) -> tuple[str, str, bytes]:
    """Decode XML while preserving the original BOM/encoding for write-back."""
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8"), "utf-8", b"\xef\xbb\xbf"
    if data.startswith(b"\xff\xfe"):
        return data[2:].decode("utf-16-le"), "utf-16-le", b"\xff\xfe"
    if data.startswith(b"\xfe\xff"):
        return data[2:].decode("utf-16-be"), "utf-16-be", b"\xfe\xff"

    declaration = data[:512]
    match = re.search(br"encoding\s*=\s*['\"]([^'\"]+)['\"]", declaration, re.IGNORECASE)
    encoding = match.group(1).decode("ascii") if match else "utf-8"
    return data.decode(encoding), encoding, b""


def encode_xml_text(text: str, encoding: str, bom: bytes) -> bytes:
    return bom + text.encode(encoding)


def local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag.rsplit(":", 1)[-1]


def validate_and_list_datasets(xml_text: str) -> list[str]:
    """Validate XML and return every DataSet/@Name in document order."""
    root = ET.fromstring(xml_text)
    names: list[str] = []
    for element in root.iter():
        if local_name(element.tag).lower() == "dataset":
            name = element.attrib.get("Name") or element.attrib.get("name")
            if name:
                names.append(name)
    return names


def xml_escape_text_value(value: str) -> str:
    """Escape only characters that are unsafe in XML element text."""
    return value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# -----------------------------------------------------------------------------
# RDL modification
# -----------------------------------------------------------------------------
def replace_snowflake_source(
    command_text: str,
    new_server: str,
    new_warehouse: str,
) -> tuple[str, list[tuple[str, str, str, str, str, str]]]:
    """
    Replace Source = Snowflake.Databases(server, warehouse, ...).

    Returns:
        modified_command_text,
        list of tuples:
          before_server, after_server, before_warehouse, after_warehouse,
          before_source_code, after_source_code
    """
    changes: list[tuple[str, str, str, str, str, str]] = []

    def replacement(match: re.Match[str]) -> str:
        before_server = match.group("server")
        before_warehouse = match.group("warehouse")

        # In raw XML text, element content still needs XML escaping.
        server = xml_escape_text_value(new_server)
        warehouse = xml_escape_text_value(new_warehouse)

        after = (
            f'{match.group("head")}'
            f'{match.group("q1")}{server}{match.group("q1")}'
            f'{match.group("separator")}'
            f'{match.group("q2")}{warehouse}{match.group("q2")}'
        )

        changes.append(
            (
                before_server,
                new_server,
                before_warehouse,
                new_warehouse,
                match.group(0),
                after,
            )
        )
        return after

    modified = SNOWFLAKE_SOURCE_RE.sub(replacement, command_text)
    return modified, changes


def modify_rdl_text(
    xml_text: str,
    new_server: str,
    new_warehouse: str,
    input_file: Path,
    output_file: Path,
) -> tuple[str, list[AuditRow], int]:
    """Modify Snowflake Source expressions dataset-by-dataset without reserializing XML."""
    audit_rows: list[AuditRow] = []
    total_changes = 0
    now = datetime.now().astimezone().isoformat(timespec="seconds")

    parts: list[str] = []
    cursor = 0

    for dataset_match in DATASET_BLOCK_RE.finditer(xml_text):
        parts.append(xml_text[cursor:dataset_match.start()])
        dataset_block = dataset_match.group(0)
        attrs = dataset_match.group("attrs")
        name_match = DATASET_NAME_RE.search(attrs)
        dataset_name = name_match.group("name") if name_match else "<unnamed>"

        dataset_changes: list[tuple[str, str, str, str, str, str]] = []

        def command_replacement(command_match: re.Match[str]) -> str:
            nonlocal dataset_changes
            command_text = command_match.group("cmd")
            modified_cmd, changes = replace_snowflake_source(
                command_text, new_server, new_warehouse
            )
            dataset_changes.extend(changes)
            return f'{command_match.group("open")}{modified_cmd}{command_match.group("close")}'

        modified_block = COMMAND_TEXT_RE.sub(command_replacement, dataset_block)
        parts.append(modified_block)
        cursor = dataset_match.end()

        if dataset_changes:
            total_changes += len(dataset_changes)
            for change_index, change in enumerate(dataset_changes, start=1):
                (
                    before_server,
                    after_server,
                    before_warehouse,
                    after_warehouse,
                    before_code,
                    after_code,
                ) = change
                status = "MODIFIED" if len(dataset_changes) == 1 else f"MODIFIED #{change_index}"
                audit_rows.append(
                    AuditRow(
                        timestamp=now,
                        file_name=input_file.name,
                        dataset_name=dataset_name,
                        status=status,
                        before_server=before_server,
                        after_server=after_server,
                        before_warehouse=before_warehouse,
                        after_warehouse=after_warehouse,
                        before_source_code=before_code,
                        after_source_code=after_code,
                        input_path=str(input_file.resolve()),
                        output_path=str(output_file.resolve()),
                    )
                )
        else:
            audit_rows.append(
                AuditRow(
                    timestamp=now,
                    file_name=input_file.name,
                    dataset_name=dataset_name,
                    status="NO_SNOWFLAKE_SOURCE_MATCH",
                    input_path=str(input_file.resolve()),
                    output_path=str(output_file.resolve()),
                )
            )

    parts.append(xml_text[cursor:])
    return "".join(parts), audit_rows, total_changes


# -----------------------------------------------------------------------------
# Excel audit log
# -----------------------------------------------------------------------------
def excel_safe(value: object) -> object:
    """Prevent strings from accidentally being interpreted as Excel formulas."""
    if isinstance(value, str) and value and value[0] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def save_audit_rows_atomic(log_path: Path, rows: Iterable[AuditRow]) -> None:
    rows = list(rows)
    if not rows:
        return

    log_path.parent.mkdir(parents=True, exist_ok=True)

    if log_path.exists():
        workbook = load_workbook(log_path)
        sheet = workbook["ModificationLog"] if "ModificationLog" in workbook.sheetnames else workbook.create_sheet("ModificationLog")
        if sheet.max_row == 1 and sheet.cell(1, 1).value is None:
            sheet.append(AUDIT_HEADERS)
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "ModificationLog"
        sheet.append(AUDIT_HEADERS)

    for row in rows:
        sheet.append([excel_safe(v) for v in row.as_list()])

    # Professional/basic audit formatting.
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    widths = {
        1: 24, 2: 30, 3: 30, 4: 26,
        5: 34, 6: 34, 7: 24, 8: 24,
        9: 55, 10: 55, 11: 45, 12: 45,
        13: 66, 14: 66,
    }
    for col_idx, width in widths.items():
        sheet.column_dimensions[get_column_letter(col_idx)].width = width

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Atomic save: write a complete temporary XLSX first, then replace.
    fd, temp_name = tempfile.mkstemp(
        prefix=log_path.stem + "_",
        suffix=".xlsx",
        dir=str(log_path.parent),
    )
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        workbook.save(temp_path)
        os.replace(temp_path, log_path)
    finally:
        workbook.close()
        if temp_path.exists():
            temp_path.unlink(missing_ok=True)


# -----------------------------------------------------------------------------
# File processing
# -----------------------------------------------------------------------------
def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def write_atomic(path: Path, data: bytes, overwrite: bool) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists() and not overwrite:
        raise FileExistsError(
            f"Output file already exists: {path}. Use --overwrite to replace it."
        )

    fd, temp_name = tempfile.mkstemp(
        prefix=path.stem + "_",
        suffix=path.suffix,
        dir=str(path.parent),
    )
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        temp_path.write_bytes(data)
        os.replace(temp_path, path)
    finally:
        if temp_path.exists():
            temp_path.unlink(missing_ok=True)


def process_rdl_file(
    input_file: Path,
    output_dir: Path,
    log_path: Path,
    new_server: str,
    new_warehouse: str,
    overwrite: bool = False,
) -> FileProcessResult:
    output_file = output_dir / input_file.name
    original_bytes = input_file.read_bytes()
    original_hash = sha256_bytes(original_bytes)
    xml_text, encoding, bom = decode_xml_bytes(original_bytes)

    # XML validation + authoritative dataset list.
    dataset_names = validate_and_list_datasets(xml_text)

    print(f"\nFile: {input_file.name}")
    print("DataSets found:")
    if dataset_names:
        for index, name in enumerate(dataset_names, start=1):
            print(f"  {index}. {name}")
    else:
        print("  (none)")

    modified_text, audit_rows, changed_count = modify_rdl_text(
        xml_text=xml_text,
        new_server=new_server,
        new_warehouse=new_warehouse,
        input_file=input_file,
        output_file=output_file,
    )

    # If there is no valid target, keep the original in place.
    if changed_count == 0:
        for row in audit_rows:
            row.original_sha256 = original_hash
        save_audit_rows_atomic(log_path, audit_rows)
        print("  No matching Source = Snowflake.Databases(server, warehouse, ...) was found.")
        print("  Original file was NOT deleted.")
        return FileProcessResult(
            input_file=input_file,
            output_file=None,
            dataset_names=dataset_names,
            changed_count=0,
            deleted_original=False,
        )

    modified_bytes = encode_xml_text(modified_text, encoding, bom)
    modified_hash = sha256_bytes(modified_bytes)

    # Validate the modified XML before writing anything permanent.
    validate_and_list_datasets(modified_text)

    for row in audit_rows:
        row.original_sha256 = original_hash
        row.modified_sha256 = modified_hash

    output_written = False
    try:
        # 1. Write modified RDL atomically.
        write_atomic(output_file, modified_bytes, overwrite=overwrite)
        output_written = True

        # 2. Write audit log atomically.
        save_audit_rows_atomic(log_path, audit_rows)

        # 3. Only after both succeed, remove the original.
        input_file.unlink()
    except Exception:
        # Keep source file safe. If this run created an output but did not finish,
        # remove the incomplete transaction's output.
        if output_written and output_file.exists():
            output_file.unlink(missing_ok=True)
        raise

    print(f"  Modified Snowflake Source occurrence(s): {changed_count}")
    print(f"  Output: {output_file}")
    print(f"  Audit log: {log_path}")
    print("  Original input file deleted after successful output + audit save.")

    return FileProcessResult(
        input_file=input_file,
        output_file=output_file,
        dataset_names=dataset_names,
        changed_count=changed_count,
        deleted_original=True,
    )


def load_config(path: Path) -> dict:
    if not path.exists():
        return {}
    with path.open("r", encoding="utf-8") as handle:
        data = json.load(handle)
    if not isinstance(data, dict):
        raise ValueError("config.json must contain a JSON object.")
    return data


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Update Source = Snowflake.Databases(server, warehouse, ...) in Power BI RDL files, "
            "write modified files to an output folder, and maintain an Excel audit log."
        )
    )
    parser.add_argument("--config", default="config.json", help="Path to JSON configuration file.")
    parser.add_argument("--input-dir", help="Folder containing incoming .rdl files.")
    parser.add_argument("--output-dir", help="Folder for modified .rdl files.")
    parser.add_argument("--audit-file", help="Excel audit log path (.xlsx).")
    parser.add_argument("--server", help="New Snowflake connection server name.")
    parser.add_argument("--warehouse", help="New Snowflake warehouse name.")
    parser.add_argument("--overwrite", action="store_true", help="Allow replacing an existing output RDL.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    config = load_config(Path(args.config))

    input_dir = Path(args.input_dir or config.get("input_dir", "input"))
    output_dir = Path(args.output_dir or config.get("output_dir", "output"))
    audit_file = Path(args.audit_file or config.get("audit_file", "logs/rdl_modification_log.xlsx"))
    server = args.server or config.get("snowflake_server")
    warehouse = args.warehouse or config.get("snowflake_warehouse")
    overwrite = bool(args.overwrite or config.get("overwrite", False))

    if not server:
        print("ERROR: Snowflake server is required. Set snowflake_server in config.json or use --server.", file=sys.stderr)
        return 2
    if not warehouse:
        print("ERROR: Snowflake warehouse is required. Set snowflake_warehouse in config.json or use --warehouse.", file=sys.stderr)
        return 2

    input_dir.mkdir(parents=True, exist_ok=True)
    output_dir.mkdir(parents=True, exist_ok=True)
    audit_file.parent.mkdir(parents=True, exist_ok=True)

    rdl_files = sorted(input_dir.glob("*.rdl"))
    if not rdl_files:
        print(f"No .rdl files found in: {input_dir.resolve()}")
        return 0

    failures = 0
    print(f"Found {len(rdl_files)} RDL file(s).")
    print(f"Target Snowflake server: {server}")
    print(f"Target Snowflake warehouse: {warehouse}")

    for rdl_file in rdl_files:
        try:
            process_rdl_file(
                input_file=rdl_file,
                output_dir=output_dir,
                log_path=audit_file,
                new_server=server,
                new_warehouse=warehouse,
                overwrite=overwrite,
            )
        except Exception as exc:
            failures += 1
            print(f"\nERROR processing {rdl_file.name}: {exc}", file=sys.stderr)
            print("Original file was kept in the input folder.", file=sys.stderr)

    if failures:
        print(f"\nCompleted with {failures} failure(s).", file=sys.stderr)
        return 1

    print("\nAll RDL files processed successfully.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
